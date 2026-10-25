# Raport xlsx din work/comparatie.json:  python -I tools/raport_fise_plan.py work/comparatie.json RAPORT.xlsx
# Prima filă = rezumat cu ce e de făcut; celelalte au 2–5 coloane și au înghețat doar rândul de antet.
import sys, json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

d = json.load(open(sys.argv[1], encoding='utf-8'))
F = 'Georgia'
GRI = PatternFill('solid', fgColor='E7E6E6')
s = Side(style='thin', color='BFBFBF'); BORD = Border(left=s, right=s, top=s, bottom=s)
SUS = Alignment(vertical='top', wrap_text=True)

def tabel(ws, cap, randuri, latimi, note=()):
    for j, t in enumerate(cap, 1):
        c = ws.cell(row=1, column=j, value=t)
        c.font = Font(name=F, size=10, bold=True, color='000000'); c.fill = GRI; c.border = BORD
        c.alignment = Alignment(vertical='center')
    for i, r in enumerate(randuri, 2):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=F, size=10); c.border = BORD; c.alignment = SUS
    for j, w in enumerate(latimi, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'
    r = len(randuri) + 3
    for t in note:
        c = ws.cell(row=r, column=1, value=t); c.font = Font(name=F, size=9, italic=True, color='595959'); r += 1

def eticheta(o):                                   # „ARA 4105 · Teoria Structurilor III”, fără traducerea din paranteză
    cod, _, den = o['cheie'].partition('|')
    den = re.sub(r'\s*\([^)]*\)', '', den).split(' / ')[0].strip()
    return f"{cod.strip()} · {den}"

def nr(v):
    v = float(v or 0)
    return int(v) if v.is_integer() else v

def ore(x):
    h = [nr(v) for v in x['h']]
    return f"{nr(x['cred'])} cr · C{h[0]} S{h[1]} L{h[2]} P{h[3]}"

def cat(o):
    if not o['plan']: return 'nelegata'
    if any(x.split(' ')[0] in ('credite', 'C', 'S', 'L', 'P') for x in o['dif']): return 'ore'
    if o['dif']: return 'alte'
    return 'ok'

out = sorted(d['out'], key=lambda o: ((o['plan'] or o['fisa'])['sem'] or 0, o['cheie']))
grup = {k: [o for o in out if cat(o) == k] for k in ('ore', 'alte', 'nelegata', 'ok')}

wb = Workbook()
rz = wb.active; rz.title = 'Rezumat'
tabel(rz, ['Situație', 'Fișe', 'Ce e de făcut'], [
    ['Ore sau credite greșite în fișă', len(grup['ore']), 'Se corectează singure la „Actualizează FD”. Lista: fila „Ore și credite”.'],
    ['Diferă doar semestrul sau forma de evaluare', len(grup['alte']), 'Tot de la planul vechi; se corectează la „Actualizează FD”. Lista: fila „Alte diferențe”.'],
    ['Fișe care nu se leagă de plan', len(grup['nelegata']), 'Decizie manuală pentru fiecare: fila „Nelegate”.'],
    ['Discipline din plan fără fișă', len(d['faraFisa']), 'Fișe de creat sau de legat: fila „Plan fără fișă”.'],
    ['Fișe corecte', len(grup['ok']), 'Nimic.'],
], [44, 7, 80], note=[
    f'Comparație între cele {len(out)} de fișe și planul oficial din mamă (INFO → PLAN_FOAIE_CURENTA = „PLAN ACTUAL”, 2026/2027, {d["nPlan"]} discipline).',
    'Cauza diferențelor: datele tehnice din fișe au fost copiate din foaia PLAN veche a mamei, care are erori.',
    'Pași în mamă: 1. Preia disciplinele din PLAN → 2. Generează URL/stare fișe → 3. Actualizează FD. Apoi se descarcă fișele și se reface comparația.',
])

ws = wb.create_sheet('Ore și credite')
tabel(ws, ['Disciplina (din fișă)', 'Sem.', 'În fișă acum', 'În plan (corect)', 'Ce se schimbă'],
      [[eticheta(o), o['plan']['sem'], ore(o['fisa']), ore(o['plan']),
        '; '.join(x for x in o['dif'] if x.split(' ')[0] in ('credite', 'C', 'S', 'L', 'P'))] for o in grup['ore']],
      [46, 6, 22, 22, 30], note=['cr = credite; C/S/L/P = ore pe săptămână de curs / seminar / laborator / proiect.'])

ws = wb.create_sheet('Alte diferențe')
tabel(ws, ['Disciplina (din fișă)', 'Ce diferă (fișă → plan)'],
      [[eticheta(o), '; '.join(o['dif'])] for o in grup['alte']], [52, 60],
      note=['Semestrul: fișele vechi aveau semestrul din an (1–2), planul are semestrul din program (1–12).',
            'Forma de evaluare: planul nou folosește E (examen) și V (verificare); fișele au și C (colocviu), P (proiect).'])

def de_facut(o):
    if 'Limb' in o['fisier']:
        return 'Planul are un singur rând „Limbi străine”, iar fișele sunt două (EN și FR). De decis: două rânduri în DISCIPLINE sau o singură fișă.'
    if 'Religious' in o['fisier']:
        return 'Titlu în engleză; probabil „Arhitectura spațiului eclezial” (sem. 8). De ales manual în foaia NELEGATE.'
    return 'Nu mai e în planul 2026/2027. De arhivat sau de legat manual în NELEGATE.'
ws = wb.create_sheet('Nelegate')
tabel(ws, ['Fișa', 'Ce e de făcut'], [[o['cheie'].replace(' | ', ' · '), de_facut(o)] for o in grup['nelegata']], [50, 90])

PED = re.compile(r'pedagog|psihologia educa|didactica|managementul clasei|instruire asistat|examen de absolvire|comunicare intercultural', re.I)
ws = wb.create_sheet('Plan fără fișă')
tabel(ws, ['Disciplina din plan', 'Sem.', 'Observație'],
      [[f"{r['cod'] or '(fără cod)'} · {r['den']}", r['sem'],
        'Modulul psihopedagogic' if PED.search(r['den']) else 'Fișă de creat, sau o fișă veche cu alt nume'] for r in d['faraFisa']],
      [62, 6, 40])

wb.save(sys.argv[2])
print('salvat:', {k: len(v) for k, v in grup.items()}, 'fara fisa:', len(d['faraFisa']))
