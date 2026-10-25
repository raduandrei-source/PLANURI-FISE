# Lista nouă de competențe (ARACIS 1–8 × C/A/R, cu referințe D46) și tabelul cod vechi → cod nou.
#   python -I tools/competente_noi.py <mama.xlsx> <ri_aracis.json> <iesire.xlsx> <mapare.json>
# ri_aracis.json: [[n, [cunoștințe], [aptitudini], [responsabilitate], discipline, pe_scurt, "litere D46"], ...]
import sys, re, json, difflib, unicodedata
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

mama, ri_json, iesire, mapare_json = sys.argv[1:5]
RI = json.load(open(ri_json, encoding='utf-8'))
CAT = {'C': 'Cunoștințe', 'A': 'Aptitudini', 'R': 'Responsabilitate și autonomie'}
VECHI_CAT = {'CN': 'C', 'AP': 'A', 'RA': 'R'}

def norm(s):
    s = unicodedata.normalize('NFD', str(s or '').lower())
    s = re.sub(r'[̀-ͯ]', '', s)
    return re.sub(r'[^a-z0-9]', '', s)
def fara_subiect(t):
    return re.sub(r'^\s*Studentul\s*/\s*absolventul\s*:?\s*', '', str(t or ''), flags=re.I)

# 1. lista nouă: C1..C8, A1..A8, R1..R8
noi, propozitii = [], []          # propozitii: (cod nou, text normalizat al unei propoziții ARACIS)
for n, cn, ap, ra, disc, scurt, d46 in RI:
    ref = f"(ARACIS-{n}, " + ', '.join(f"D46-{l.strip()}" for l in d46.split(',')) + ')'
    for lit, lst in (('C', cn), ('A', ap), ('R', ra)):
        cod = f'{lit}{n}'
        noi.append([cod, CAT[lit], ' '.join(lst) + ' ' + ref, f'ARACIS-{n}: {scurt}'])
        propozitii += [(cod, norm(fara_subiect(t))) for t in lst]

# decizii pe text, verificate de mână (2026-10-25): variante adaptate ale unui rând ARACIS; liste de discipline puse greșit ca RÎ
ADAPTATE = {'AP13': ['A1'], 'RA14': ['R4'], 'RA15': ['R4']}
RENUNTA = {'RA11', 'RA13'}

# 2. codurile vechi → coduri noi (propozițiile ARACIS conținute în textul vechi sau foarte asemănătoare)
vechi = [r for r in load_workbook(mama, data_only=True)['COMPETENTE'].iter_rows(min_row=2, values_only=True)
         if r[0] and re.match(r'^(CN|AP|RA)\d+$', str(r[0]))]
mapare, randuri_map = {}, []
for cod, _cat, text, *_ in vechi:
    t = norm(fara_subiect(text))
    gasite = [c for c, p in propozitii if len(p) > 20 and (p in t or (len(t) > 30 and t in p))]
    cum = 'text ARACIS'
    if not gasite and t:
        best = max(propozitii, key=lambda cp: difflib.SequenceMatcher(None, t, cp[1]).ratio())
        if difflib.SequenceMatcher(None, t, best[1]).ratio() >= 0.85: gasite, cum = [best[0]], 'text ARACIS, cu mici diferențe'
    if cod in ADAPTATE: gasite, cum = ADAPTATE[cod], 'variantă adaptată a textului ARACIS'
    gasite = list(dict.fromkeys(gasite))
    renunta = cod in RENUNTA
    mapare[cod] = {'noi': gasite, 'text': str(text or ''), 'renunta': renunta}
    if renunta: nou, cum = '— se renunță', 'listă de discipline, nu e rezultat al învățării'
    elif gasite: nou = ', '.join(gasite)
    else: nou, cum = '→ Specificări manuale', 'text propriu al disciplinei'
    randuri_map.append([cod, nou, cum, str(text or '')])
json.dump(mapare, open(mapare_json, 'w', encoding='utf-8'), ensure_ascii=False, indent=0)

# 3. fișierul pentru mamă: foile COMPETENTE (nouă) și COMPETENTE_VECHI (tabelul de trecere)
F = 'Georgia'; GRI = PatternFill('solid', fgColor='E7E6E6')
s = Side(style='thin', color='BFBFBF'); BORD = Border(left=s, right=s, top=s, bottom=s)
def foaie(ws, cap, rows, lat):
    for j, v in enumerate(cap, 1):
        c = ws.cell(row=1, column=j, value=v); c.font = Font(name=F, size=10, bold=True, color='000000'); c.fill = GRI; c.border = BORD
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v); c.font = Font(name=F, size=10, bold=(j == 1)); c.border = BORD
            c.alignment = Alignment(vertical='top', wrap_text=True)
    for j, w in enumerate(lat, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'
wb = Workbook()
foaie(wb.active, ['COD', 'CATEGORIE', 'TEXT (Rezultatele învățării)', 'OBSERVAȚII'], noi, [7, 16, 100, 40])
wb.active.title = 'COMPETENTE'
foaie(wb.create_sheet('COMPETENTE_VECHI'), ['COD VECHI', 'COD NOU', 'CUM', 'TEXT VECHI'], randuri_map, [10, 14, 28, 100])
wb.save(iesire)
nr = sum(1 for v in mapare.values() if v['noi'])
print(f'lista nouă: {len(noi)} coduri | coduri vechi: {len(mapare)} → ARACIS: {nr}, texte proprii: {len(mapare) - nr}')
