import sys, json, glob, os
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook
src, mama, out = sys.argv[1], sys.argv[2], sys.argv[3]
fise = []
for f in sorted(glob.glob(os.path.join(src, '*.xlsx'))):
    wb = load_workbook(f, data_only=False, read_only=True)
    dt = wb['_DATE'] if '_DATE' in wb.sheetnames else wb['FISA']
    v = {}
    for row in dt.iter_rows(min_row=3, max_row=500, min_col=15, max_col=38):
        for c in row:
            if c.value not in (None, ''): v[(c.row, c.column)] = c.value
    cheie = str(v.get((10, 15), '') or '').strip()
    tab = []
    for r in range(4, 501):
        rr = [v.get((r, c), '') for c in range(24, 39)]
        if any(x not in ('', None) for x in rr): tab.append(rr)
    own = next((x for x in tab if str(x[14]).strip() == cheie), None)
    fise.append({'fisier': os.path.basename(f), 'cod': v.get((4, 15), ''), 'cheie': cheie,
                 'randuri_tabel': len(tab), 'rand': own, 'foi': wb.sheetnames})
    wb.close()
wb = load_workbook(mama, data_only=True)
ws = wb['PLAN ACTUAL']
plan = [[c if c is not None else '' for c in row] for row in ws.iter_rows(values_only=True)]
ws = wb['DISCIPLINE']
disc = [[c if c is not None else '' for c in row[:25]] for row in ws.iter_rows(min_row=2, values_only=True) if row[1]]
json.dump({'fise': fise, 'plan': plan, 'disc': disc}, open(out, 'w', encoding='utf-8'), ensure_ascii=False, default=str)
print('fise', len(fise), 'fara rand propriu', sum(1 for x in fise if not x['rand']), 'disc', len(disc))
from collections import Counter
print('randuri in tabel:', Counter(x['randuri_tabel'] for x in fise).most_common(5))
